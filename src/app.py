"""
app.py
หน้าเว็บ Streamlit สำหรับ News Perspective Comparator
รัน: streamlit run app.py
"""

import html
import io
import os
import re
import sys
import time

import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

try:
    for _k, _v in st.secrets.items():
        if isinstance(_v, str):
            os.environ.setdefault(_k, _v)
except Exception:
    pass

import news_comparator as nc  # noqa: E402

MAX_RESULTS = 8
DAYS_BACK = 30
MIN_LETTERS = 3
MAX_INPUT_CHARS = 200
SENTIMENT_COLORS = {"positive": "#4ade80", "negative": "#f87171", "neutral": "#facc15"}

EXCEL_COLUMNS = [
    ("Source", "source", 22),
    ("Country", "country", 10),
    ("Title", "title", 50),
    ("URL", "url", 40),
    ("Sentiment", "sentiment", 12),
    ("Sentiment score", "sentiment_score", 15),
    ("Summary", "summary", 60),
    ("Angle", "angle", 50),
    ("Unique point", "unique_point", 50),
]

st.set_page_config(page_title="News Perspective Comparator", layout="wide")
st.title("News Perspective Comparator")


def build_excel(keyword: str, results: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    ws.append([f"Topic: {keyword}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([name for name, _, _ in EXCEL_COLUMNS])
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    for r in results:
        ws.append([r.get(key, "") for _, key, _ in EXCEL_COLUMNS])
        fill = SENTIMENT_COLORS.get(r.get("sentiment", ""))
        if fill:
            sentiment_cell = ws.cell(row=ws.max_row, column=5)
            sentiment_cell.fill = PatternFill("solid", fgColor=fill.lstrip("#"))
    for i, (_, _, width) in enumerate(EXCEL_COLUMNS, 1):
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = width
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def validate_topic(text: str) -> str | None:
    if not text:
        return "Please enter a news topic."
    if re.search(r"[^\x00-\x7F‘’“”–—]", text):
        return "Please type the topic in English only (no Thai text, emoji or special symbols)."
    if len(re.findall(r"[A-Za-z]", text)) < MIN_LETTERS:
        return "The topic is too short. Please describe the news in a few English words, e.g. \"Tesla deliveries\"."
    return None


def gemini_error_message(err: nc.errors.APIError) -> str:
    if err.code in (400, 401, 403):
        return f"Gemini rejected the request (error {err.code}). Please check that GEMINI_API_KEY is correct."
    if err.code == 429:
        return "Gemini quota is used up. Please wait a few minutes and try again."
    return f"Gemini is busy or temporarily down (error {err.code}). Please try again in a moment."


@st.cache_data(show_spinner=False, ttl=3600)
def get_keyword(text: str) -> str:
    return nc.clean_keyword(text)


@st.cache_data(show_spinner=False, ttl=3600)
def get_articles(keyword: str, user_input: str) -> tuple[list[dict], str]:
    return nc.find_articles(keyword, user_input, MAX_RESULTS, DAYS_BACK)


@st.cache_data(show_spinner=False, ttl=3600)
def judge_article(topic: str, title: str, text: str) -> tuple[str, str]:
    time.sleep(nc.JUDGE_DELAY)
    return nc.judge_relevance(topic, title, text)


@st.cache_data(show_spinner=False, ttl=3600)
def get_results(keyword: str, relevant: list[dict]) -> list[dict]:
    return nc.compare_perspectives(keyword, relevant)


text = st.text_input(
    "Enter a news topic you are interested in (English)",
    placeholder="e.g. Brydon Carse charged over nightclub incident",
    max_chars=MAX_INPUT_CHARS,
)

if st.button("Analyze", type="primary"):
    text = text.strip()
    st.session_state.pop("analysis", None)
    problem = validate_topic(text)
    if problem:
        st.warning(problem)
        st.stop()
    bar = st.progress(0.0, text="Extracting search keywords...")
    try:
        keyword = get_keyword(text)
        bar.progress(0.1, text="Searching for news articles...")

        articles, keyword = get_articles(keyword, text)
        bar.progress(0.2, text=f"Checking relevance (0/{len(articles)})...")
        if not articles:
            bar.empty()
            st.warning("No news found. Try a different keyword.")
            st.stop()

        relevant = []
        for i, a in enumerate(articles, 1):
            label, reason = judge_article(text, a["title"], a["text"][:nc.JUDGE_CHARS])
            print(f"  {'✓' if label == nc.POS else '✗'} [{a['source']}] {a['title'][:70]}  ({reason})")
            if label == nc.POS:
                relevant.append(a)
            bar.progress(0.2 + 0.65 * i / len(articles), text=f"Checking relevance ({i}/{len(articles)})...")
        if not relevant:
            bar.empty()
            st.warning("No relevant news found. Try a different keyword.")
            st.stop()

        bar.progress(0.85, text="Comparing perspectives...")
        results = get_results(keyword, relevant)
        bar.progress(1.0, text="Done")
        bar.empty()
    except nc.errors.APIError as err:
        bar.empty()
        st.error(gemini_error_message(err))
        st.stop()
    except (RuntimeError, SystemExit) as err:
        bar.empty()
        st.error(str(err))
        st.stop()
    except ValueError as err:
        bar.empty()
        st.error("Gemini returned an answer in an unexpected format. Please click Analyze again.")
        with st.expander("Error details"):
            st.exception(err)
        st.stop()
    except Exception as err:
        bar.empty()
        st.error("Something unexpected went wrong. Please try again.")
        with st.expander("Error details"):
            st.exception(err)
        st.stop()

    st.session_state["analysis"] = {"keyword": keyword, "results": results}

analysis = st.session_state.get("analysis")
if analysis:
    keyword, results = analysis["keyword"], analysis["results"]
    slug = re.sub(r"[^a-z0-9]+", "_", keyword.lower()).strip("_") or "result"

    st.subheader(f"Topic: {keyword}")
    head_left, head_right = st.columns([3, 1], vertical_alignment="center")
    head_left.write(f"Found perspectives from {len(results)} sources")
    head_right.download_button(
        "Export to Excel",
        data=build_excel(keyword, results),
        file_name=f"{slug}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )

    cols = st.columns(2)
    for i, r in enumerate(results):
        with cols[i % 2].container(border=True):
            sentiment = r.get("sentiment", "neutral")
            bg = SENTIMENT_COLORS.get(sentiment, SENTIMENT_COLORS["neutral"])
            score = r.get("sentiment_score")
            badge = f"{sentiment}" if isinstance(score, (int, float)) else sentiment
            st.markdown(
                f"<div style='display:flex; justify-content:space-between; align-items:center; gap:12px;'>"
                f"<h3 style='margin:0; padding:0;'>{html.escape(r['source'])}</h3>"
                f"<span style='background:{bg}; color:#000; border-radius:6px; padding:2px 10px; "
                f"font-size:13px; font-weight:600; white-space:nowrap;'>{html.escape(badge)}</span>"
                f"</div>",
                unsafe_allow_html=True,
            )
            st.markdown(f"[{r['title']}]({r['url']})")
            for label, key in (("Summary", "summary"), ("Angle", "angle"), ("Unique point", "unique_point")):
                st.markdown(
                    f"<div style='background:rgba(128,128,128,0.12); border-radius:8px; margin-bottom:10px;"
                    f"padding:10px 14px; margin-top:8px;'>"
                    f"<div style='font-size:12px; font-weight:700; letter-spacing:0.06em; "
                    f"text-transform:uppercase; opacity:0.65; margin-bottom:4px;'>{label}</div>"
                    f"<div>{html.escape(r.get(key, ''))}</div></div>",
                    unsafe_allow_html=True,
                )
